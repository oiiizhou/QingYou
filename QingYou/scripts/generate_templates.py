"""可选开发工具：重新生成仓库内的有效空白 Office 模板。"""
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
import xlwt
import subprocess

root = Path(__file__).resolve().parent.parent / "Resources" / "Templates"
root.mkdir(parents=True, exist_ok=True)
document = Document()
document.core_properties.author = ""
document.core_properties.title = ""
document.save(root / "Blank.docx")
workbook = Workbook()
workbook.active.title = "Sheet1"
workbook.properties.creator = ""
workbook.save(root / "Blank.xlsx")
presentation = Presentation()
presentation.slides.add_slide(presentation.slide_layouts[6])
presentation.core_properties.author = ""
presentation.save(root / "Blank.pptx")
subprocess.run(["textutil", "-convert", "doc", str(root / "Blank.docx"), "-output", str(root / "Blank.doc")], check=True)
legacy_book = xlwt.Workbook()
legacy_book.add_sheet("Sheet1")
legacy_book.save(str(root / "Blank.xls"))
# Blank.ppt 使用本机演示软件将 Blank.pptx 另存为 PPT 97–2003 后提交。
# 保留仓库中的真实二进制模板，不能通过重命名生成。
print("已生成有效的空白 Office 模板。")
